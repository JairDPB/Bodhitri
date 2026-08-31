#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Extrae las imagenes de producto de un Excel y las renombra con el ID del producto.

Funciona con CUALQUIER archivo .xlsx: no hay rutas, hojas ni columnas fijas en el
codigo. Todo se localiza por NOMBRE de encabezado y se puede sobreescribir por
linea de comandos.

    Excel con imagenes  ->  <ID>.png, <ID>.png, ...  ->  imagenes.zip

Ese ZIP es el que se sube a Business Central con la pagina "Importar Imagenes de
Productos (ZIP)", que asigna cada imagen al producto cuyo N.º coincide con el
nombre del archivo.

USO
---
    python extraer_imagenes_excel.py libro.xlsx
    python extraer_imagenes_excel.py libro.xlsx --col-id REFERENCIA --salida ./out
    python extraer_imagenes_excel.py libro.xlsx --hojas BOLDY,MATCH --dry-run

    # Recuperar las imagenes de mayor resolucion cuando el libro las tenga
    # "en celda" (requiere Pillow):
    python extraer_imagenes_excel.py libro.xlsx --alta-resolucion

COMO ENCUENTRA LAS COSAS
------------------------
1. Recorre las hojas del libro (todas, o las que se indiquen con --hojas).
2. En cada hoja busca la fila de encabezado: la primera que contenga alguno de
   los nombres de columna de ID (--col-id, o los alias por defecto).
3. Localiza tambien la columna de imagen (--col-imagen o alias por defecto).
4. Lee los anclajes de los dibujos (xl/drawings): cada imagen flotante declara
   la celda sobre la que esta colocada.
5. Con la fila del anclaje, lee el ID de esa fila y renombra la imagen.

Requiere: openpyxl.  Opcional: Pillow (solo para --alta-resolucion y para medir).
"""

from __future__ import annotations

import argparse
import re
import shutil
import sys
import unicodedata
import xml.etree.ElementTree as ET
import zipfile
from collections import defaultdict
from pathlib import Path

try:
    import openpyxl
except ImportError:
    sys.exit("Falta openpyxl. Instalalo con:  pip install openpyxl")


# --- Espacios de nombres de OOXML -------------------------------------------
NS = {
    "main": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
    "rel":  "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "pr":   "http://schemas.openxmlformats.org/package/2006/relationships",
    "xdr":  "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing",
    "a":    "http://schemas.openxmlformats.org/drawingml/2006/main",
}
R_EMBED = f"{{{NS['rel']}}}embed"
R_ID = f"{{{NS['rel']}}}id"

# Nombres de columna aceptados por defecto. El ID es lo que se usa para renombrar.
ALIAS_ID = ["DYNAMICS", "DINAMYCS", "DINAMICS", "DYNAMYCS",
            "ID", "CODIGO", "COD", "REFERENCIA", "REF", "SKU",
            "ITEM", "ARTICULO", "PRODUCTO", "NO.", "N."]
ALIAS_IMAGEN = ["IMAGEN", "IMAGENES", "IMAGE", "FOTO", "FOTOGRAFIA", "PHOTO", "PICTURE"]

EXT_VALIDAS = {".png", ".jpg", ".jpeg", ".gif", ".bmp"}
# Longitud del campo "No." de la tabla Item en Business Central.
MAX_ID_BC = 20


# =============================================================================
# UTILIDADES
# =============================================================================

def normalizar(texto) -> str:
    """Mayusculas, sin tildes, sin espacios repetidos. Para comparar encabezados."""
    if texto is None:
        return ""
    s = str(texto).strip().upper()
    s = unicodedata.normalize("NFKD", s)
    s = "".join(c for c in s if not unicodedata.combining(c))
    return re.sub(r"\s+", " ", s)


def nombre_archivo_valido(texto: str) -> bool:
    """Un ID solo sirve si puede ser nombre de archivo y N.º de producto en BC."""
    if not texto or len(texto) > MAX_ID_BC:
        return False
    return not re.search(r'[<>:"/\\|?*\x00-\x1f]', texto)


def medir_png(datos: bytes):
    """Ancho y alto de un PNG leyendo la cabecera IHDR, sin depender de Pillow."""
    if len(datos) < 24 or datos[:8] != b"\x89PNG\r\n\x1a\n":
        return None
    return (int.from_bytes(datos[16:20], "big"), int.from_bytes(datos[20:24], "big"))


def medir(datos: bytes, ext: str):
    m = medir_png(datos)
    if m:
        return m
    try:                                    # otros formatos: solo si hay Pillow
        from PIL import Image
        import io
        with Image.open(io.BytesIO(datos)) as im:
            return im.size
    except Exception:
        return None


# =============================================================================
# LECTURA DE LA ESTRUCTURA DEL XLSX
# =============================================================================

def leer_rels(z: zipfile.ZipFile, ruta_parte: str) -> dict:
    """Devuelve {rId: destino_absoluto} del archivo .rels que acompana a una parte."""
    p = Path(ruta_parte)
    ruta_rels = f"{p.parent.as_posix()}/_rels/{p.name}.rels"
    if ruta_rels not in z.namelist():
        return {}
    raiz = ET.fromstring(z.read(ruta_rels))
    out = {}
    for r in raiz.findall("pr:Relationship", NS):
        destino = r.get("Target", "")
        if destino.startswith("/"):
            destino = destino[1:]
        else:                                # resolver la ruta relativa
            destino = (p.parent / destino).as_posix()
            while "/../" in destino:
                destino = re.sub(r"[^/]+/\.\./", "", destino, count=1)
        out[r.get("Id")] = destino
    return out


def hojas_del_libro(z: zipfile.ZipFile) -> dict:
    """{nombre_de_hoja: ruta de su sheetN.xml}"""
    raiz = ET.fromstring(z.read("xl/workbook.xml"))
    rels = leer_rels(z, "xl/workbook.xml")
    out = {}
    for h in raiz.findall(".//main:sheets/main:sheet", NS):
        rid = h.get(R_ID)
        if rid in rels:
            out[h.get("name")] = rels[rid]
    return out


def anclajes_de_hoja(z: zipfile.ZipFile, ruta_hoja: str) -> list:
    """
    Devuelve [(columna, fila, ruta_de_la_imagen)] de las imagenes flotantes.

    Las coordenadas del anclaje son BASE 0: la fila 0 del XML es la fila 1 de Excel.
    """
    rels_hoja = leer_rels(z, ruta_hoja)
    dibujos = [d for d in rels_hoja.values() if "drawings/drawing" in d]
    salida = []

    for dib in dibujos:
        if dib not in z.namelist():
            continue
        raiz = ET.fromstring(z.read(dib))
        rels_dib = leer_rels(z, dib)

        for etiqueta in ("twoCellAnchor", "oneCellAnchor"):
            for anc in raiz.findall(f"xdr:{etiqueta}", NS):
                desde = anc.find("xdr:from", NS)
                if desde is None:
                    continue
                col_el, fila_el = desde.find("xdr:col", NS), desde.find("xdr:row", NS)
                if col_el is None or fila_el is None:
                    continue

                blip = anc.find(".//xdr:pic/xdr:blipFill/a:blip", NS)
                if blip is None:
                    continue
                destino = rels_dib.get(blip.get(R_EMBED))
                if not destino or destino not in z.namelist():
                    continue

                salida.append((int(col_el.text) + 1,   # a base 1
                               int(fila_el.text) + 1,
                               destino))
    return salida


# =============================================================================
# LOCALIZAR ENCABEZADO Y COLUMNAS
# =============================================================================

def localizar_columnas(ws, alias_id, alias_img, fila_encabezado=None, filas_a_mirar=25):
    """
    Busca la fila de encabezado y las columnas de ID e imagen, por NOMBRE.

    Devuelve (fila_encabezado, columna_id, columna_imagen) o None si no hay ID.
    """
    filas = ([fila_encabezado] if fila_encabezado
             else range(1, min(ws.max_row, filas_a_mirar) + 1))

    for f in filas:
        col_id = col_img = None
        for c in range(1, min(ws.max_column, 60) + 1):
            n = normalizar(ws.cell(f, c).value)
            if not n:
                continue
            if col_id is None and n in alias_id:
                col_id = c
            if col_img is None and n in alias_img:
                col_img = c
        if col_id:
            return f, col_id, col_img
    return None


# =============================================================================
# PROCESO PRINCIPAL
# =============================================================================

def procesar(args) -> int:
    origen = Path(args.excel)
    if not origen.exists():
        print(f"No existe el archivo: {origen}", file=sys.stderr)
        return 2

    salida = Path(args.salida)
    dir_img = salida / "imagenes"
    if not args.dry_run:
        dir_img.mkdir(parents=True, exist_ok=True)

    alias_id = [normalizar(x) for x in (args.col_id.split(",") if args.col_id else ALIAS_ID)]
    alias_img = [normalizar(x) for x in (args.col_imagen.split(",") if args.col_imagen else ALIAS_IMAGEN)]

    print(f"Leyendo {origen.name}")
    z = zipfile.ZipFile(origen)
    wb = openpyxl.load_workbook(origen, data_only=True, read_only=False)
    rutas_hoja = hojas_del_libro(z)

    hojas_pedidas = [h.strip() for h in args.hojas.split(",")] if args.hojas else None

    exportadas = {}          # id -> dict con los datos
    huerfanas, avisos = [], []
    total_anclajes = 0

    for ws in wb.worksheets:
        if hojas_pedidas and ws.title not in hojas_pedidas:
            continue
        if ws.title not in rutas_hoja or ws.max_row < 2:
            continue

        cols = localizar_columnas(ws, alias_id, alias_img, args.fila_encabezado)
        if not cols:
            avisos.append(f"{ws.title}: sin columna de ID reconocible; hoja omitida.")
            continue
        fila_cab, col_id, col_img = cols

        anclajes = anclajes_de_hoja(z, rutas_hoja[ws.title])
        total_anclajes += len(anclajes)
        print(f"  {ws.title:<14} encabezado fila {fila_cab} | ID col {col_id} | "
              f"imagen col {col_img if col_img else '-'} | {len(anclajes)} imagenes")

        for col, fila, parte in anclajes:
            # Las imagenes decorativas (logos, fotos de ambiente) suelen estar
            # ancladas fuera de la columna de imagen. Si sabemos cual es, se exige
            # que el anclaje caiga ahi; si no lo sabemos, se acepta cualquiera.
            if col_img and col != col_img:
                huerfanas.append(f"{ws.title}!fila {fila}: imagen en columna {col}, "
                                 f"no en la columna de imagen ({col_img})")
                continue
            if fila <= fila_cab:
                huerfanas.append(f"{ws.title}!fila {fila}: por encima del encabezado")
                continue

            crudo = ws.cell(fila, col_id).value
            ident = str(crudo).strip() if crudo is not None else ""
            if not ident:
                huerfanas.append(f"{ws.title}!fila {fila}: la fila no tiene ID")
                continue
            if not nombre_archivo_valido(ident):
                avisos.append(f"{ws.title}!fila {fila}: ID '{ident}' no sirve como "
                              f"nombre de archivo o supera {MAX_ID_BC} caracteres; omitido.")
                continue

            datos = z.read(parte)
            ext = Path(parte).suffix.lower()
            if ext not in EXT_VALIDAS:
                avisos.append(f"{ws.title}!fila {fila}: formato {ext} no soportado.")
                continue

            if ident in exportadas:
                avisos.append(f"ID duplicado '{ident}': ya venia de "
                              f"{exportadas[ident]['hoja']}!{exportadas[ident]['fila']}, "
                              f"ahora en {ws.title}!{fila}. Se conserva el primero.")
                continue

            exportadas[ident] = {
                "hoja": ws.title, "fila": fila, "parte": parte,
                "bytes": len(datos), "medida": medir(datos, ext), "ext": ext,
                "datos": datos,
            }

    # --- Recuperar las imagenes de mayor resolucion, si se pidio -------------
    if args.alta_resolucion:
        recuperadas = recuperar_alta_resolucion(z, exportadas, avisos)
        if recuperadas:
            print(f"  Alta resolucion: {recuperadas} imagenes sustituidas por su version grande")

    # --- Escribir ------------------------------------------------------------
    if args.dry_run:
        print("\n(--dry-run: no se escribio nada)")
    else:
        for ident, info in exportadas.items():
            (dir_img / f"{ident}{info['ext']}").write_bytes(info["datos"])

        ruta_zip = salida / args.zip
        with zipfile.ZipFile(ruta_zip, "w", zipfile.ZIP_DEFLATED) as zf:
            for ident, info in exportadas.items():
                # Sin carpetas dentro del ZIP: Business Central toma el nombre
                # del archivo como N.º de producto.
                zf.writestr(f"{ident}{info['ext']}", info["datos"])

        escribir_reporte(salida / "reporte.md", origen, exportadas, huerfanas,
                         avisos, total_anclajes, wb, rutas_hoja, alias_id,
                         alias_img, args)

    # --- Resumen -------------------------------------------------------------
    print(f"\n  Imagenes ancladas encontradas : {total_anclajes}")
    print(f"  Exportadas con ID              : {len(exportadas)}")
    print(f"  Descartadas (decorativas, etc.): {len(huerfanas)}")
    print(f"  Avisos                         : {len(avisos)}")
    if not args.dry_run:
        print(f"\n  {salida.resolve()}")
        print(f"    imagenes/  ({len(exportadas)} archivos)")
        print(f"    {args.zip}")
        print(f"    reporte.md")
    for a in avisos[:10]:
        print(f"  ! {a}")
    if len(avisos) > 10:
        print(f"  ! ... y {len(avisos) - 10} avisos mas (ver reporte.md)")

    return 0 if exportadas else 1


def recuperar_alta_resolucion(z, exportadas, avisos) -> int:
    """
    Algunos libros guardan DOS copias de cada foto: la flotante (pequena, la que
    tiene el anclaje) y otra "en celda" de mas resolucion, bajo xl/richData.

    La cadena que ata la celda con la imagen grande (metadata -> richValue) se
    pierde con facilidad al guardar el libro. Cuando eso pasa, las imagenes
    grandes quedan huerfanas dentro del paquete y no hay forma de saber a que
    producto pertenecen leyendo el formato.

    Aqui se recuperan por PARECIDO: cada imagen grande se compara con las
    pequenas ya mapeadas y se queda con su gemela. Es una inferencia, no un dato
    del archivo: por eso solo se aplica con --alta-resolucion y solo se acepta
    un emparejamiento cuando la similitud es muy alta y ademas es mutua.

    Requiere Pillow. Si no esta, se avisa y se deja todo como esta.
    """
    try:
        from PIL import Image
    except ImportError:
        avisos.append("--alta-resolucion necesita Pillow (pip install Pillow). "
                      "Se conservan las imagenes flotantes.")
        return 0

    import io

    usadas = {info["parte"] for info in exportadas.values()}
    candidatas = [n for n in z.namelist()
                  if n.startswith("xl/media/") and n not in usadas
                  and Path(n).suffix.lower() in EXT_VALIDAS]
    if not candidatas:
        return 0

    def huella(datos, lado=48):
        """Miniatura normalizada en color, conservando la relacion de aspecto."""
        try:
            with Image.open(io.BytesIO(datos)) as im:
                im = im.convert("RGB")
                lienzo = Image.new("RGB", (lado, lado), (255, 255, 255))
                im.thumbnail((lado, lado), Image.LANCZOS)
                lienzo.paste(im, ((lado - im.width) // 2, (lado - im.height) // 2))
                return list(lienzo.getdata())
        except Exception:
            return None

    def similitud(a, b):
        """Correlacion de Pearson sobre los tres canales."""
        pa = [v for px in a for v in px]
        pb = [v for px in b for v in px]
        n = len(pa)
        ma, mb = sum(pa) / n, sum(pb) / n
        num = sum((x - ma) * (y - mb) for x, y in zip(pa, pb))
        da = sum((x - ma) ** 2 for x in pa) ** 0.5
        db = sum((y - mb) ** 2 for y in pb) ** 0.5
        return num / (da * db) if da and db else 0.0

    hg = {c: huella(z.read(c)) for c in candidatas}
    hg = {k: v for k, v in hg.items() if v}
    hp = {i: huella(info["datos"]) for i, info in exportadas.items()}
    hp = {k: v for k, v in hp.items() if v}

    # Mejor candidata para cada producto, y viceversa: solo vale si es mutua.
    mejor_por_id, mejor_por_cand = {}, defaultdict(lambda: (None, -1))
    for ident, h in hp.items():
        puntuaciones = sorted(((similitud(h, hg[c]), c) for c in hg), reverse=True)
        if puntuaciones:
            s, c = puntuaciones[0]
            mejor_por_id[ident] = (c, s)
            if s > mejor_por_cand[c][1]:
                mejor_por_cand[c] = (ident, s)

    sustituidas = 0
    for ident, (cand, s) in mejor_por_id.items():
        if s < 0.97:
            continue
        if mejor_por_cand[cand][0] != ident:      # el emparejamiento no es mutuo
            avisos.append(f"'{ident}': la imagen grande mas parecida ya la reclama "
                          f"otro producto; se conserva la pequena.")
            continue
        datos = z.read(cand)
        grande, pequena = medir(datos, Path(cand).suffix), exportadas[ident]["medida"]
        if grande and pequena and (grande[0] * grande[1]) <= (pequena[0] * pequena[1]):
            continue                              # no aporta resolucion
        exportadas[ident].update({"datos": datos, "parte": cand,
                                  "bytes": len(datos), "medida": grande,
                                  "ext": Path(cand).suffix.lower(),
                                  "alta_resolucion": True, "similitud": round(s, 5)})
        sustituidas += 1

    return sustituidas


def escribir_reporte(destino, origen, exportadas, huerfanas, avisos,
                     total_anclajes, wb, rutas_hoja, alias_id, alias_img, args):
    L = [f"# Extracción de imágenes — {origen.name}\n"]
    L.append(f"- Imágenes ancladas encontradas: **{total_anclajes}**")
    L.append(f"- Exportadas con ID: **{len(exportadas)}**")
    L.append(f"- Descartadas: **{len(huerfanas)}**")
    L.append(f"- Avisos: **{len(avisos)}**\n")

    por_hoja = defaultdict(int)
    for info in exportadas.values():
        por_hoja[info["hoja"]] += 1
    if por_hoja:
        L.append("## Por hoja\n")
        L.append("| Hoja | Imágenes |")
        L.append("|---|---:|")
        for h, n in sorted(por_hoja.items()):
            L.append(f"| {h} | {n} |")
        L.append("")

    if exportadas:
        medidas = [i["medida"] for i in exportadas.values() if i["medida"]]
        if medidas:
            lados = sorted(max(m) for m in medidas)
            L.append("## Resolución (lado largo, px)\n")
            L.append(f"mín {lados[0]} · mediana {lados[len(lados)//2]} · máx {lados[-1]}\n")

        L.append("## Imágenes exportadas\n")
        L.append("| ID | Origen | Tamaño | px | Alta resol. |")
        L.append("|---|---|---:|---|:-:|")
        for ident in sorted(exportadas):
            i = exportadas[ident]
            px = f"{i['medida'][0]}×{i['medida'][1]}" if i["medida"] else "?"
            alta = "sí" if i.get("alta_resolucion") else ""
            L.append(f"| `{ident}` | {i['hoja']}!{i['fila']} | {i['bytes']:,} B | {px} | {alta} |")
        L.append("")

    if huerfanas:
        L.append("## Descartadas\n")
        L.append("Imágenes ancladas que no corresponden a un producto "
                 "(decorativas, logos, fotos de ambiente):\n")
        for h in huerfanas:
            L.append(f"- {h}")
        L.append("")

    if avisos:
        L.append("## Avisos\n")
        for a in avisos:
            L.append(f"- {a}")
        L.append("")

    destino.write_text("\n".join(L), encoding="utf-8")


# =============================================================================
# MAIN
# =============================================================================

def main() -> int:
    ap = argparse.ArgumentParser(
        description="Extrae las imágenes de un Excel renombradas con el ID del producto "
                    "y las empaqueta en un ZIP para Business Central.",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog="Ejemplo:\n"
               "  python extraer_imagenes_excel.py libro.xlsx --col-id REFERENCIA\n")
    ap.add_argument("excel", help="Archivo .xlsx de origen")
    ap.add_argument("--salida", default="./salida", help="Carpeta de salida (por defecto ./salida)")
    ap.add_argument("--zip", default="imagenes.zip", help="Nombre del ZIP (por defecto imagenes.zip)")
    ap.add_argument("--col-id", default=None,
                    help="Nombre(s) de la columna con el ID del producto, separados por coma. "
                         f"Por defecto prueba: {', '.join(ALIAS_ID[:6])}...")
    ap.add_argument("--col-imagen", default=None,
                    help="Nombre(s) de la columna donde están las imágenes. Si se indica, se "
                         "descartan las imágenes ancladas en otras columnas (logos, ambientes).")
    ap.add_argument("--hojas", default=None, help="Hojas a procesar, separadas por coma")
    ap.add_argument("--fila-encabezado", type=int, default=None,
                    help="Fila del encabezado, si no se quiere detectar automáticamente")
    ap.add_argument("--alta-resolucion", action="store_true",
                    help="Intenta recuperar las copias de mayor resolución que algunos libros "
                         "guardan como imagen 'en celda'. Requiere Pillow. Ver la nota del código: "
                         "el emparejamiento es por parecido, no un dato del archivo.")
    ap.add_argument("--dry-run", action="store_true", help="Solo informa, no escribe nada")
    args = ap.parse_args()

    try:
        return procesar(args)
    except zipfile.BadZipFile:
        print("El archivo no es un .xlsx válido (¿es .xls antiguo?).", file=sys.stderr)
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
